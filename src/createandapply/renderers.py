"""ATS-safe output: single column, standard headings, real text, no tables/images/text boxes."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
from xml.sax.saxutils import escape

FONT = "Calibri"


def _dates(job: dict) -> str:
    return f"{job.get('start', '')} – {job.get('end', 'Present')}".strip(" –")


def _contact_line(c: dict) -> str:
    return " | ".join(x for x in [c.get("location"), c.get("phone"), c.get("email"),
                                  c.get("linkedin"), c.get("github")] if x)


# ---------------------------------------------------------------- DOCX
def _docx_base() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.8)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(2)
    return doc


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11.5)


def resume_docx(r: dict, path: Path) -> None:
    doc = _docx_base()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(r["name"])
    run.bold = True
    run.font.size = Pt(18)
    if r.get("headline"):
        h = doc.add_paragraph(r["headline"])
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c = doc.add_paragraph(_contact_line(r["contact"]))
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if r.get("summary"):
        _heading(doc, "Professional Summary")
        doc.add_paragraph(r["summary"])
    if r.get("skills"):
        _heading(doc, "Skills")
        for cat, items in r["skills"].items():
            p = doc.add_paragraph()
            p.add_run(f"{cat}: ").bold = True
            p.add_run(", ".join(items))
    _heading(doc, "Professional Experience")
    for job in r["experience"]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.add_run(job["title"]).bold = True
        p.add_run(f" | {job['company']}" + (f" | {job['location']}" if job.get("location") else ""))
        p.add_run(f" | {_dates(job)}")
        for b in job.get("bullets", []):
            doc.add_paragraph(b, style="List Bullet")
    if r.get("education"):
        _heading(doc, "Education")
        for e in r["education"]:
            doc.add_paragraph(", ".join(filter(None, [e.get("degree"), e.get("school"), str(e.get("year", ""))])))
    if r.get("certifications"):
        _heading(doc, "Certifications")
        for cert in r["certifications"]:
            doc.add_paragraph(str(cert), style="List Bullet")
    doc.core_properties.author = r["name"]
    doc.core_properties.title = f"{r['name']} - Resume"
    doc.save(path)


def letter_docx(text: str, name: str, path: Path) -> None:
    doc = _docx_base()
    for block in text.split("\n\n"):
        lines = block.split("\n")
        if all(l.startswith("- ") for l in lines[1:]) and len(lines) > 1 and not lines[0].startswith("- "):
            doc.add_paragraph(lines[0])
            for l in lines[1:]:
                doc.add_paragraph(l[2:], style="List Bullet")
        else:
            p = doc.add_paragraph("\n".join(lines))
            p.paragraph_format.space_after = Pt(8)
    doc.core_properties.author = name
    doc.core_properties.title = f"{name} - Cover Letter"
    doc.save(path)


# ---------------------------------------------------------------- PDF (text-based, selectable)
def _styles():
    ss = getSampleStyleSheet()
    base = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=10, leading=13)
    return {
        "base": base,
        "name": ParagraphStyle("n", parent=base, fontName="Helvetica-Bold", fontSize=17, alignment=1, leading=20),
        "center": ParagraphStyle("c", parent=base, alignment=1),
        "h": ParagraphStyle("h", parent=base, fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=2),
    }


def _pdf(path: Path, story: list, title: str, author: str) -> None:
    SimpleDocTemplate(str(path), pagesize=letter, leftMargin=0.8 * inch, rightMargin=0.8 * inch,
                      topMargin=0.7 * inch, bottomMargin=0.7 * inch, title=title, author=author).build(story)


def resume_pdf(r: dict, path: Path) -> None:
    s = _styles()
    e = escape
    story = [Paragraph(e(r["name"]), s["name"])]
    if r.get("headline"):
        story.append(Paragraph(e(r["headline"]), s["center"]))
    story.append(Paragraph(e(_contact_line(r["contact"])), s["center"]))
    if r.get("summary"):
        story += [Paragraph("PROFESSIONAL SUMMARY", s["h"]), Paragraph(e(r["summary"]), s["base"])]
    if r.get("skills"):
        story.append(Paragraph("SKILLS", s["h"]))
        for cat, items in r["skills"].items():
            story.append(Paragraph(f"<b>{e(cat)}:</b> {e(', '.join(items))}", s["base"]))
    story.append(Paragraph("PROFESSIONAL EXPERIENCE", s["h"]))
    for job in r["experience"]:
        loc = f" | {job['location']}" if job.get("location") else ""
        story.append(Paragraph(f"<b>{e(job['title'])}</b> | {e(job['company'])}{e(loc)} | {e(_dates(job))}", s["base"]))
        story.append(ListFlowable([ListItem(Paragraph(e(b), s["base"])) for b in job.get("bullets", [])],
                                  bulletType="bullet", start="•", leftIndent=14))
        story.append(Spacer(1, 4))
    if r.get("education"):
        story.append(Paragraph("EDUCATION", s["h"]))
        for ed in r["education"]:
            story.append(Paragraph(e(", ".join(filter(None, [ed.get("degree"), ed.get("school"), str(ed.get("year", ""))]))), s["base"]))
    if r.get("certifications"):
        story.append(Paragraph("CERTIFICATIONS", s["h"]))
        story.append(ListFlowable([ListItem(Paragraph(e(str(c)), s["base"])) for c in r["certifications"]],
                                  bulletType="bullet", start="•", leftIndent=14))
    _pdf(path, story, f"{r['name']} - Resume", r["name"])


def letter_pdf(text: str, name: str, path: Path) -> None:
    s = _styles()
    story = []
    for block in text.split("\n\n"):
        lines = block.split("\n")
        if len(lines) > 1 and all(l.startswith("- ") for l in lines[1:]) and not lines[0].startswith("- "):
            story.append(Paragraph(escape(lines[0]), s["base"]))
            story.append(ListFlowable([ListItem(Paragraph(escape(l[2:]), s["base"])) for l in lines[1:]],
                                      bulletType="bullet", start="•", leftIndent=14))
        else:
            story.append(Paragraph(escape(block).replace("\n", "<br/>"), s["base"]))
        story.append(Spacer(1, 8))
    _pdf(path, story, f"{name} - Cover Letter", name)
